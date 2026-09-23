"""ExportService — rapports (PDF, DOCX, Markdown, TXT, JSON), sous-titres (SRT, VTT), chapitres et clips."""
from __future__ import annotations

import io
import json
import time
from pathlib import Path

from backend import config, db
from backend.services import analysis_store, ffmpeg_service, storage_service, transcription_service
from backend.services.util import fmt_srt_time, fmt_time

PINK = "#E83967"
DARK = "#2D2D2D"
CLOUDY = "#B1ADA1"
SOFT = "#F2D5D0"
PAMPAS = "#F4F3EE"

REPORT_FORMATS = {"pdf", "docx", "md", "txt", "json", "srt", "vtt"}
CATEGORY_LABELS = {"décision": "Décision", "action": "Action", "information": "Information", "question": "Question",
                   "idée": "Idée", "opportunité": "Opportunité", "risque": "Risque", "citation": "Citation",
                   "date": "Date", "chiffre": "Chiffre"}
STATUS_LABELS = {"todo": "À faire", "doing": "En cours", "done": "Terminé"}
_SYMBOLS = {"★": "*", "☆": "-", "✓": "v", "☐": "[ ]", "⚠": "!"}


def _logo_png() -> Path | None:
    for name in ("logo.png", "favicon.png"):
        p = config.FRONTEND_DIR / "img" / name
        if p.exists():
            return p
    return None


def stars(importance: float) -> str:
    n = max(1, min(5, round((importance or 0) * 5)))
    return "★" * n + "☆" * (5 - n)


# --------------------------------------------------------------------------- collecte

def collect(audio_id: int, sections: set[str] | None = None, start: float | None = None,
            end: float | None = None) -> dict:
    audio = db.one("SELECT * FROM audio_files WHERE id = ?", (audio_id,))
    if not audio:
        raise KeyError("Audio introuvable")

    def in_range(t):
        return t is None or ((start is None or t >= start) and (end is None or t <= end))

    summ = analysis_store.current(audio_id, "summary")
    segs = [s for s in transcription_service.segments(audio_id) if in_range(s["start"])]
    data = {
        "audio": audio,
        "summary": summ["content"] if summ else {},
        "summary_versions": len(analysis_store.history(audio_id, "summary")),
        "extra": {k: (analysis_store.current(audio_id, k) or {}).get("content", {}).get("markdown", "")
                  for k in ("summary_detailed", "summary_chronological", "summary_thematic")},
        "briefing": (analysis_store.current(audio_id, "briefing") or {}).get("content", {}),
        "chapters": [c for c in db.all("SELECT * FROM chapters WHERE audio_id = ? ORDER BY start", (audio_id,))
                     if in_range(c["start"])],
        "highlights": [h for h in db.all("SELECT * FROM highlights WHERE audio_id = ? ORDER BY importance DESC, start",
                                         (audio_id,)) if in_range(h["start"])],
        "decisions": [r for r in db.all("SELECT * FROM decisions WHERE audio_id = ? ORDER BY time", (audio_id,))
                      if in_range(r["time"])],
        "actions": [r for r in db.all("SELECT * FROM actions WHERE audio_id = ? ORDER BY time", (audio_id,))
                    if in_range(r["time"])],
        "questions": [r for r in db.all("SELECT * FROM questions WHERE audio_id = ? ORDER BY time", (audio_id,))
                      if in_range(r["time"])],
        "risks": [r for r in db.all("SELECT * FROM risks WHERE audio_id = ? ORDER BY time", (audio_id,))
                  if in_range(r["time"])],
        "people": db.all("SELECT * FROM people WHERE audio_id = ? ORDER BY mentions DESC", (audio_id,)),
        "topics": db.all("SELECT * FROM topics WHERE audio_id = ? ORDER BY duration DESC", (audio_id,)),
        "segments": segs,
        "stats": stats(audio_id),
    }
    for p in data["people"]:
        p["topics"] = db.jloads(p["topics"], [])
    for c in data["chapters"]:
        c["topics"] = db.jloads(c["topics"], [])
    if sections is not None:
        for key in list(data):
            if key not in sections and key not in ("audio", "stats"):
                data[key] = [] if isinstance(data[key], list) else ({} if isinstance(data[key], dict) else data[key])
    return data


def stats(audio_id: int) -> dict:
    audio = db.one("SELECT duration FROM audio_files WHERE id = ?", (audio_id,)) or {"duration": 0}
    segs = transcription_service.segments(audio_id)
    words = sum(len(s["text"].split()) for s in segs)
    speech = sum(max(0, s["end"] - s["start"]) for s in segs)
    count = lambda t: db.scalar(f"SELECT COUNT(*) FROM {t} WHERE audio_id = ?", (audio_id,)) or 0  # noqa: E731
    speakers = db.all("SELECT label, name, (SELECT COALESCE(SUM(end - start), 0) FROM transcription_segments s "
                      "WHERE s.speaker_id = sp.id) AS talk FROM speakers sp WHERE audio_id = ?", (audio_id,))
    return {
        "duration": audio["duration"], "words": words, "segments": len(segs), "speech_time": round(speech, 1),
        "words_per_minute": round(words / (speech / 60), 1) if speech > 30 else None,
        "speakers": len(speakers), "speaker_times": speakers,
        "chapters": count("chapters"), "highlights": count("highlights"), "actions": count("actions"),
        "decisions": count("decisions"), "questions": count("questions"), "risks": count("risks"),
        "people": count("people"), "topics": count("topics"), "bookmarks": count("bookmarks"),
        "annotations": count("annotations"), "clips": count("clips"),
        "avg_confidence": round(sum((s["confidence"] or 0) for s in segs) / len(segs), 3) if segs else None,
        "edited_segments": sum(1 for s in segs if s["edited"]),
    }


# --------------------------------------------------------------------------- formats texte

def to_srt(segments: list[dict], offset: float = 0) -> str:
    out = []
    for i, s in enumerate(segments, 1):
        out.append(f"{i}\n{fmt_srt_time(s['start'] - offset)} --> {fmt_srt_time(s['end'] - offset)}\n{s['text']}\n")
    return "\n".join(out)


def to_vtt(segments: list[dict], offset: float = 0) -> str:
    out = ["WEBVTT", ""]
    for s in segments:
        out.append(f"{fmt_srt_time(s['start'] - offset, '.')} --> {fmt_srt_time(s['end'] - offset, '.')}")
        out.append(s["text"])
        out.append("")
    return "\n".join(out)


def to_markdown(d: dict, include_transcript: bool = True) -> str:
    a, s, st = d["audio"], d["summary"], d["stats"]
    L = [f"# {s.get('title') or a['title']}", ""]
    meta = [f"**Durée :** {fmt_time(a['duration'])}", f"**Date :** {(a.get('recorded_at') or a['created_at'])[:10]}"]
    if a.get("participants"):
        meta.append(f"**Participants :** {a['participants']}")
    L += [" · ".join(meta), "", "> Document généré localement par BONJOUR IA — Audio Intelligence Workspace.", ""]
    if s.get("tldr"):
        L += ["## TL;DR", "", s["tldr"], ""]
    if s.get("short"):
        L += ["## Résumé", "", s["short"], ""]
    if s.get("executive"):
        L += ["## Résumé exécutif", "", s["executive"], ""]
    for key, title in (("summary_detailed", "Résumé détaillé"), ("summary_chronological", "Résumé chronologique"),
                       ("summary_thematic", "Résumé par thème")):
        if d["extra"].get(key):
            L += [f"## {title}", "", d["extra"][key], ""]
    if s.get("key_points"):
        L += ["## À retenir", ""] + [f"- {x}" for x in s["key_points"]] + [""]
    if d["chapters"]:
        L += ["## Chapitres", ""]
        for c in d["chapters"]:
            L.append(f"- **{fmt_time(c['start'])} — {c['title']}**" + (f" : {c['summary']}" if c["summary"] else ""))
        L.append("")
    if d["highlights"]:
        L += ["## Moments clés", ""]
        for h in d["highlights"][:40]:
            L.append(f"- `{fmt_time(h['start'])}` {stars(h['importance'])} *{CATEGORY_LABELS.get(h['category'], h['category'])}* — {h['text']}")
        L.append("")
    if d["decisions"]:
        L += ["## Décisions", ""] + [f"- ✓ {r['text']} (`{fmt_time(r['time'])}`)" for r in d["decisions"]] + [""]
    if d["actions"]:
        L += ["## Actions", ""]
        for r in d["actions"]:
            box = "x" if r["status"] == "done" else " "
            extra = ", ".join(x for x in (f"Responsable : {r['owner']}" if r["owner"] else "",
                                          f"Échéance : {r['deadline']}" if r["deadline"] else "",
                                          f"Source : {fmt_time(r['time'])}") if x)
            L.append(f"- [{box}] {r['text']} — {extra}")
        L.append("")
    if d["questions"]:
        L += ["## Questions ouvertes", ""] + [f"- ❓ {r['text']} (`{fmt_time(r['time'])}`)" for r in d["questions"]] + [""]
    if d["risks"]:
        L += ["## Risques", ""]
        for r in d["risks"]:
            tag = "Interprétation IA" if r["kind"] == "inference" else "Mentionné"
            L.append(f"- ⚠ {r['text']} — *{tag}* (`{fmt_time(r['time'])}`)")
        L.append("")
    if s.get("problems"):
        L += ["## Problèmes", ""] + [f"- {x}" for x in s["problems"]] + [""]
    if s.get("figures"):
        L += ["## Chiffres", ""] + [f"- {x}" for x in s["figures"]] + [""]
    if s.get("dates"):
        L += ["## Dates", ""] + [f"- {x}" for x in s["dates"]] + [""]
    if d["people"]:
        L += ["## Personnes", ""]
        for p in d["people"]:
            L.append(f"- **{p['name']}**" + (f" — {p['role']}" if p["role"] else "") + f" ({p['mentions']} mentions)")
        L.append("")
    if d["topics"]:
        L += ["## Sujets", ""] + [f"- {t['name']} — {t['occurrences']} occurrences, {fmt_time(t['duration'])}"
                                  for t in d["topics"]] + [""]
    L += ["## Statistiques", "", f"- Durée : {fmt_time(st['duration'])}", f"- Mots : {st['words']}",
          f"- Segments : {st['segments']}", f"- Chapitres : {st['chapters']}", f"- Highlights : {st['highlights']}",
          f"- Actions : {st['actions']}", f"- Décisions : {st['decisions']}", f"- Questions : {st['questions']}"]
    if st.get("words_per_minute"):
        L.append(f"- Débit de parole : {st['words_per_minute']} mots/min")
    L.append("")
    if include_transcript and d["segments"]:
        L += ["## Transcription", ""]
        for seg in d["segments"]:
            spk = f"**{seg['speaker_name'] or seg['speaker']}** " if seg.get("speaker") else ""
            L.append(f"`{fmt_time(seg['start'])}` {spk}{seg['text']}  ")
    return "\n".join(L).strip() + "\n"


def to_txt(d: dict) -> str:
    import re

    md = to_markdown(d)
    md = re.sub(r"[*`>]", "", md)
    md = re.sub(r"^#+\s*", "", md, flags=re.M)
    return md


def to_json(d: dict) -> str:
    out = dict(d)
    out["exported_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
    out["generator"] = "BONJOUR IA — Audio Intelligence Workspace"
    return json.dumps(out, ensure_ascii=False, indent=2, default=str)


# --------------------------------------------------------------------------- DOCX

def to_docx(d: dict, include_transcript: bool = True) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    pink = RGBColor(0xE8, 0x39, 0x67)
    dark = RGBColor(0x2D, 0x2D, 0x2D)
    cloudy = RGBColor(0xB1, 0xAD, 0xA1)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(11)
    st.font.color.rgb = dark
    for name, size, color in (("Title", 26, pink), ("Heading 1", 18, pink), ("Heading 2", 14, cloudy)):
        s = doc.styles[name]
        s.font.name = "Poppins"
        s.font.size = Pt(size)
        s.font.color.rgb = color
        s.font.bold = name != "Title"
    a, s = d["audio"], d["summary"]
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2.2)
    logo = _logo_png()
    if logo:
        doc.add_picture(str(logo), height=Cm(1.6))
    footer = section.footer.paragraphs[0]
    footer.text = "BONJOUR IA — Audio Intelligence Workspace · Document généré localement"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = cloudy

    # Couverture
    doc.add_paragraph()
    kicker = doc.add_paragraph()
    r = kicker.add_run("FICHE DE SYNTHÈSE")
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = pink
    doc.add_paragraph(s.get("title") or a["title"], style="Title")
    meta = doc.add_paragraph()
    mr = meta.add_run(f"Durée {fmt_time(a['duration'])} · {(a.get('recorded_at') or a['created_at'])[:10]}"
                      + (f" · {a['participants']}" if a.get("participants") else ""))
    mr.font.color.rgb = cloudy
    mr.font.name = "Poppins"
    if s.get("tldr"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        rr = p.add_run(s["tldr"])
        rr.italic = True
        rr.font.size = Pt(13)
    doc.add_page_break()

    def h(text):
        doc.add_heading(text, level=1)

    def bullets(items):
        for it in items:
            doc.add_paragraph(it, style="List Bullet")

    if s.get("short"):
        h("Résumé")
        doc.add_paragraph(s["short"])
    if s.get("executive"):
        h("Résumé exécutif")
        doc.add_paragraph(s["executive"])
    for key, title in (("summary_detailed", "Résumé détaillé"), ("summary_chronological", "Résumé chronologique"),
                       ("summary_thematic", "Résumé par thème")):
        if d["extra"].get(key):
            h(title)
            for line in d["extra"][key].splitlines():
                t = line.strip().lstrip("#").strip()
                if not t:
                    continue
                if line.strip().startswith(("-", "*")):
                    doc.add_paragraph(t.lstrip("-* ").replace("**", ""), style="List Bullet")
                else:
                    doc.add_paragraph(t.replace("**", ""))
    if s.get("key_points"):
        h("À retenir")
        bullets(s["key_points"])
    if d["chapters"]:
        h("Chapitres")
        for c in d["chapters"]:
            p = doc.add_paragraph()
            rt = p.add_run(f"{fmt_time(c['start'])}  ")
            rt.font.color.rgb = pink
            rt.bold = True
            p.add_run(c["title"]).bold = True
            if c["summary"]:
                p.add_run(f" — {c['summary']}")
    if d["highlights"]:
        h("Moments clés")
        for hl in d["highlights"][:30]:
            p = doc.add_paragraph()
            rt = p.add_run(f"{fmt_time(hl['start'])}  {stars(hl['importance'])}  ")
            rt.font.color.rgb = pink
            p.add_run(f"{CATEGORY_LABELS.get(hl['category'], hl['category'])} — ").bold = True
            p.add_run(hl["text"])
    for key, title, fmt in (
        ("decisions", "Décisions", lambda r: f"✓ {r['text']} ({fmt_time(r['time'])})"),
        ("actions", "Actions", lambda r: f"☐ {r['text']}" + (f" — {r['owner']}" if r["owner"] else "")
         + (f" — échéance : {r['deadline']}" if r["deadline"] else "") + f" ({STATUS_LABELS.get(r['status'], r['status'])}, {fmt_time(r['time'])})"),
        ("questions", "Questions ouvertes", lambda r: f"{r['text']} ({fmt_time(r['time'])})"),
        ("risks", "Risques", lambda r: f"{r['text']} — {'interprétation IA' if r['kind'] == 'inference' else 'mentionné'} ({fmt_time(r['time'])})"),
    ):
        if d[key]:
            h(title)
            bullets([fmt(r) for r in d[key]])
    for key, title in (("problems", "Problèmes"), ("figures", "Chiffres"), ("dates", "Dates")):
        if s.get(key):
            h(title)
            bullets(s[key])
    if d["people"]:
        h("Personnes")
        bullets([f"{p['name']}" + (f" — {p['role']}" if p["role"] else "") + f" ({p['mentions']} mentions)"
                 for p in d["people"]])
    if d["topics"]:
        h("Sujets")
        bullets([f"{t['name']} — {t['occurrences']} occurrences, {fmt_time(t['duration'])}" for t in d["topics"]])
    stt = d["stats"]
    h("Statistiques")
    table = doc.add_table(rows=0, cols=2)
    for label, value in (("Durée", fmt_time(stt["duration"])), ("Mots", stt["words"]), ("Segments", stt["segments"]),
                         ("Chapitres", stt["chapters"]), ("Highlights", stt["highlights"]),
                         ("Actions", stt["actions"]), ("Décisions", stt["decisions"]),
                         ("Questions", stt["questions"])):
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(value)
    if include_transcript and d["segments"]:
        doc.add_page_break()
        h("Transcription")
        for seg in d["segments"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            rt = p.add_run(f"{fmt_time(seg['start'])}  ")
            rt.font.color.rgb = cloudy
            rt.font.size = Pt(9)
            p.add_run(seg["text"])
    for p in doc.paragraphs:
        if p.style.name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- PDF

def to_pdf(d: dict, include_transcript: bool = True) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)
    from xml.sax.saxutils import escape

    pink, dark, cloudy = colors.HexColor(PINK), colors.HexColor(DARK), colors.HexColor(CLOUDY)
    # Polices TrueType Windows si disponibles (accents et symboles), sinon polices PDF standard
    font, font_b, font_i, serif = "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Times-Roman"
    sym_font = None
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        win = Path(r"C:\Windows\Fonts")
        if (win / "segoeui.ttf").exists():
            pdfmetrics.registerFont(TTFont("Segoe", str(win / "segoeui.ttf")))
            pdfmetrics.registerFont(TTFont("Segoe-B", str(win / "segoeuib.ttf")))
            pdfmetrics.registerFont(TTFont("Segoe-I", str(win / "segoeuii.ttf")))
            font, font_b, font_i = "Segoe", "Segoe-B", "Segoe-I"
        if (win / "georgia.ttf").exists():
            pdfmetrics.registerFont(TTFont("Georgia", str(win / "georgia.ttf")))
            serif = "Georgia"
        if (win / "seguisym.ttf").exists():
            pdfmetrics.registerFont(TTFont("SegoeSym", str(win / "seguisym.ttf")))
            sym_font = "SegoeSym"
    except Exception:
        pass

    def Para(markup, style):
        # Les symboles (★ ✓ ☐ ⚠) absents des polices de texte passent par Segoe UI Symbol
        out = []
        for ch in str(markup):
            if ch in _SYMBOLS:
                out.append(f'<font name="{sym_font}">{ch}</font>' if sym_font else _SYMBOLS[ch])
            else:
                out.append(ch)
        return Paragraph("".join(out), style)

    ss = {
        "kicker": ParagraphStyle("k", fontName=font_b, fontSize=8.5, textColor=pink, leading=12, spaceAfter=6),
        "title": ParagraphStyle("t", fontName=font, fontSize=28, leading=34, textColor=dark, spaceAfter=12),
        "meta": ParagraphStyle("m", fontName=font, fontSize=10, textColor=cloudy, leading=14, spaceAfter=18),
        "tldr": ParagraphStyle("tl", fontName=font_i, fontSize=13, leading=19, textColor=dark, leftIndent=12,
                               borderPadding=(6, 0, 6, 10), borderColor=pink, borderWidth=0),
        "h1": ParagraphStyle("h1", fontName=font_b, fontSize=15, leading=20, textColor=pink, spaceBefore=16,
                             spaceAfter=8),
        "body": ParagraphStyle("b", fontName=serif, fontSize=10.5, leading=15.5, textColor=dark, spaceAfter=6,
                               alignment=TA_LEFT),
        "li": ParagraphStyle("li", fontName=serif, fontSize=10.5, leading=15, textColor=dark, leftIndent=12,
                             bulletIndent=0, spaceAfter=3),
        "seg": ParagraphStyle("s", fontName=serif, fontSize=9.5, leading=13.5, textColor=dark, spaceAfter=2),
    }

    def P(text, style="body"):
        return Para(escape(str(text)).replace("\n", "<br/>"), ss[style])

    def T(t):
        return f'<font name="{font_b}" color="{PINK}">{fmt_time(t)}</font>'

    a, s = d["audio"], d["summary"]
    story = []
    logo = _logo_png()
    if logo:
        from reportlab.lib.utils import ImageReader

        iw, ih = ImageReader(str(logo)).getSize()
        hgt = 1.6 * cm
        story.append(Image(str(logo), width=hgt * iw / ih, height=hgt, hAlign="LEFT"))
    story += [Spacer(1, 5 * cm), P("FICHE DE SYNTHÈSE · AUDIO INTELLIGENCE", "kicker"),
              P(s.get("title") or a["title"], "title"),
              P(f"Durée {fmt_time(a['duration'])}  ·  {(a.get('recorded_at') or a['created_at'])[:10]}"
                + (f"  ·  {a['participants']}" if a.get("participants") else ""), "meta")]
    if s.get("tldr"):
        story.append(P(s["tldr"], "tldr"))
    story.append(PageBreak())

    def section(title, items):
        if items:
            story.append(P(title, "h1"))
            story.extend(items)

    def md_block(text):
        items = []
        for line in text.splitlines():
            t = line.strip()
            if not t:
                continue
            clean = escape(t.lstrip("#-* ").strip()).replace("**", "")
            items.append(Para(("• " if t.startswith(("-", "*")) else "") + clean,
                                   ss["li" if t.startswith(("-", "*")) else "body"]))
        return items

    section("Résumé", [P(s["short"])] if s.get("short") else [])
    section("Résumé exécutif", [P(s["executive"])] if s.get("executive") else [])
    for key, title in (("summary_detailed", "Résumé détaillé"), ("summary_chronological", "Résumé chronologique"),
                       ("summary_thematic", "Résumé par thème")):
        section(title, md_block(d["extra"][key]) if d["extra"].get(key) else [])
    section("À retenir", [Para("• " + escape(x), ss["li"]) for x in s.get("key_points", [])])
    section("Chapitres", [Para(f"{T(c['start'])}&nbsp;&nbsp;<b>{escape(c['title'])}</b>"
                                    + (f" — {escape(c['summary'])}" if c["summary"] else ""), ss["li"])
                          for c in d["chapters"]])
    section("Moments clés", [Para(f"{T(h['start'])}&nbsp;&nbsp;<font color='{PINK}'>{stars(h['importance'])}</font>"
                                       f"&nbsp;&nbsp;<b>{escape(CATEGORY_LABELS.get(h['category'], h['category']))}</b> — "
                                       f"{escape(h['text'])}", ss["li"]) for h in d["highlights"][:30]])
    section("Décisions", [Para(f"✓ {escape(r['text'])} &nbsp;{T(r['time'])}", ss["li"]) for r in d["decisions"]])
    section("Actions", [Para(f"☐ <b>{escape(r['text'])}</b>"
                                  + (f" — Responsable : {escape(r['owner'])}" if r["owner"] else "")
                                  + (f" — Échéance : {escape(r['deadline'])}" if r["deadline"] else "")
                                  + f" — {STATUS_LABELS.get(r['status'], r['status'])} &nbsp;{T(r['time'])}", ss["li"])
                        for r in d["actions"]])
    section("Questions ouvertes", [Para(f"? {escape(r['text'])} &nbsp;{T(r['time'])}", ss["li"])
                                   for r in d["questions"]])
    section("Risques", [Para(f"⚠ {escape(r['text'])} — <i>{'Interprétation IA' if r['kind'] == 'inference' else 'Mentionné'}</i>"
                                  f" &nbsp;{T(r['time'])}", ss["li"]) for r in d["risks"]])
    for key, title in (("problems", "Problèmes"), ("figures", "Chiffres"), ("dates", "Dates")):
        section(title, [Para("• " + escape(x), ss["li"]) for x in s.get(key, [])])
    section("Personnes", [Para(f"• <b>{escape(p['name'])}</b>" + (f" — {escape(p['role'])}" if p["role"] else "")
                                    + f" ({p['mentions']} mentions)", ss["li"]) for p in d["people"]])
    section("Sujets", [Para(f"• {escape(t['name'])} — {t['occurrences']} occurrences, {fmt_time(t['duration'])}",
                                 ss["li"]) for t in d["topics"]])
    stt = d["stats"]
    rows = [["Durée", fmt_time(stt["duration"])], ["Mots", stt["words"]], ["Segments", stt["segments"]],
            ["Chapitres", stt["chapters"]], ["Highlights", stt["highlights"]], ["Actions", stt["actions"]],
            ["Décisions", stt["decisions"]], ["Questions", stt["questions"]]]
    tbl = Table([[str(x) for x in r] for r in rows], colWidths=[6 * cm, 6 * cm], hAlign="LEFT")
    tbl.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), font, 10), ("TEXTCOLOR", (0, 0), (0, -1), cloudy),
                             ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E3DC")),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    section("Statistiques", [tbl])
    if include_transcript and d["segments"]:
        story.append(PageBreak())
        story.append(P("Transcription", "h1"))
        for seg in d["segments"]:
            story.append(Para(f"<font name='{font}' color='{CLOUDY}' size='8.5'>{fmt_time(seg['start'])}</font>"
                                   f"&nbsp;&nbsp;{escape(seg['text'])}", ss["seg"]))

    def on_page(canvas, doc_):
        canvas.saveState()
        canvas.setFont(font, 7.5)
        canvas.setFillColor(cloudy)
        canvas.drawString(2 * cm, 1.2 * cm, "BONJOUR IA — Audio Intelligence Workspace · Traitement 100 % local")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"{doc_.page}")
        canvas.setStrokeColor(pink)
        canvas.setLineWidth(2)
        canvas.line(2 * cm, A4[1] - 1.2 * cm, 3.2 * cm, A4[1] - 1.2 * cm)
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title=s.get("title") or a["title"], author="BONJOUR IA")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# --------------------------------------------------------------------------- points d'entrée

MIME = {"pdf": "application/pdf", "md": "text/markdown; charset=utf-8", "txt": "text/plain; charset=utf-8",
        "json": "application/json", "srt": "application/x-subrip; charset=utf-8", "vtt": "text/vtt; charset=utf-8",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "mp3": "audio/mpeg", "wav": "audio/wav"}


def export_report(audio_id: int, fmt: str, sections: set[str] | None = None,
                  include_transcript: bool = True) -> tuple[bytes, str, str]:
    if fmt not in REPORT_FORMATS:
        raise ValueError("Format inconnu")
    d = collect(audio_id, sections)
    name = storage_service.slugify(d["summary"].get("title") or d["audio"]["title"])
    if fmt == "srt":
        body = to_srt(d["segments"]).encode("utf-8")
    elif fmt == "vtt":
        body = to_vtt(d["segments"]).encode("utf-8")
    elif fmt == "md":
        body = to_markdown(d, include_transcript).encode("utf-8")
    elif fmt == "txt":
        body = to_txt(d).encode("utf-8")
    elif fmt == "json":
        body = to_json(d).encode("utf-8")
    elif fmt == "docx":
        body = to_docx(d, include_transcript)
    else:
        body = to_pdf(d, include_transcript)
    return body, f"{name}.{fmt}", MIME[fmt]


def export_range(audio_id: int, start: float, end: float, fmt: str, title: str = "") -> tuple[bytes, str, str]:
    """Export d'un chapitre ou d'un clip : audio (mp3/wav), transcription (txt/md/srt), PDF."""
    audio = db.one("SELECT * FROM audio_files WHERE id = ?", (audio_id,))
    if not audio:
        raise KeyError("Audio introuvable")
    name = storage_service.slugify(f"{audio['title']}-{title or fmt_time(start)}")
    if fmt in ("mp3", "wav"):
        tmp = config.TEMP_DIR / f"{storage_service.new_uid()}.{fmt}"
        try:
            ffmpeg_service.export_segment(storage_service.audio_path(audio["stored_name"]), tmp, start, end, fmt)
            return tmp.read_bytes(), f"{name}.{fmt}", MIME[fmt]
        finally:
            tmp.unlink(missing_ok=True)
    segs = [s for s in transcription_service.segments(audio_id) if s["end"] > start and s["start"] < end]
    if fmt == "srt":
        return to_srt(segs, offset=start).encode("utf-8"), f"{name}.srt", MIME["srt"]
    if fmt in ("txt", "md"):
        head = f"# {title or 'Extrait'}\n\n**{audio['title']}** — {fmt_time(start)} → {fmt_time(end)}\n\n" if fmt == "md" \
            else f"{title or 'Extrait'}\n{audio['title']} — {fmt_time(start)} → {fmt_time(end)}\n\n"
        body = head + "\n".join((f"`{fmt_time(s['start'])}` " if fmt == "md" else f"[{fmt_time(s['start'])}] ")
                                + s["text"] + ("  " if fmt == "md" else "") for s in segs)
        return body.encode("utf-8"), f"{name}.{fmt}", MIME[fmt]
    if fmt == "pdf":
        d = collect(audio_id, {"chapters", "highlights", "decisions", "actions", "questions", "risks", "segments"},
                    start, end)
        d["summary"] = {"title": title or f"Extrait {fmt_time(start)} → {fmt_time(end)}"}
        chapter = db.one("SELECT summary FROM chapters WHERE audio_id = ? AND start = ?", (audio_id, start))
        if chapter and chapter["summary"]:
            d["summary"]["short"] = chapter["summary"]
        d["audio"] = {**audio, "duration": end - start}
        return to_pdf(d), f"{name}.pdf", MIME["pdf"]
    raise ValueError("Format inconnu")
